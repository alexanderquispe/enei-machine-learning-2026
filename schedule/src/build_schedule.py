"""Split the ML I / ML II teaching hours between instructor and TA, and write the Excel file.

Rules
-----
1. Each course is 30 hours: 18 h for the instructor (6 sessions) and 12 h for the
   teaching assistant (4 sessions), matching the 2025 edition. Every session is
   3 hours long (19:00-22:00).
2. The instructor teaches Python at the PUCP Diploma on Wednesday evenings and
   that calendar cannot be moved, so every Wednesday session goes to the TA.
3. If the Wednesday conflicts do not add up to the TA's 12 hours, the last
   session of the course is reassigned to the TA as well.

Inputs (schedule/data/processed/): enei_sessions_2026.csv, pucp_conflicts_2026.csv
Output (schedule/output/):         horario_peu_cd_2026_ml1_ml2.xlsx
"""

import datetime as dt
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
PROCESSED = ROOT / "data" / "processed"
OUTPUT = ROOT / "output"

INSTRUCTOR = "Alexander Quispe"
ASSISTANT = "Rodrigo Grijalba"

SESSION_HOURS = 3
INSTRUCTOR_SESSIONS = 6  # 18 hours
ASSISTANT_SESSIONS = 4   # 12 hours

COURSES = {"MACHING LEARNING I": "ML I", "MACHING LEARNING II": "ML II"}

# Course labels as they should appear in the Spanish-language deliverable
COURSE_LABELS = {
    "FUND. CIENCIA DE DATOS": "Fundamentos de la Ciencia de Datos",
    "INGENIERIA DE DATOS": "Ingeniería de Datos aplicado a la Estadística",
    "ESTIMACION ESTAD.": "Métodos de Estimación Estadística",
    "SEMINARIO": "Seminario",
    "DEEP LEARNING": "Deep Learning",
    "PROYECTO INTEGRADOR": "Proyecto Integrador",
    "Combate de Angamos": "FERIADO - Combate de Angamos",
}

WEEKDAYS_ES = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"]
MONTHS_ES = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
             "agosto", "setiembre", "octubre", "noviembre", "diciembre"]

# Palette
BLUE, BLUE_LIGHT = "1F4E79", "DDEBF7"
GREEN, GREEN_TEXT = "C6EFCE", "006100"
ORANGE, ORANGE_TEXT = "FCE4D6", "833C0C"
GREY, GREY_TEXT = "F2F2F2", "808080"
YELLOW = "FFF2CC"
BOX = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def load_sessions() -> tuple[pd.DataFrame, dict[dt.date, str]]:
    enei = pd.read_csv(PROCESSED / "enei_sessions_2026.csv", parse_dates=["date"])
    enei["date"] = enei["date"].dt.date

    pucp = pd.read_csv(PROCESSED / "pucp_conflicts_2026.csv", parse_dates=["date"])
    pucp["date"] = pucp["date"].dt.date
    conflicts = dict(zip(pucp["date"], pucp["label"]))

    sessions = enei[enei["course"].isin(COURSES)].copy()
    sessions["course"] = sessions["course"].map(COURSES)
    return sessions.sort_values("date").reset_index(drop=True), conflicts


def assign(sessions: pd.DataFrame, conflicts: dict[dt.date, str]) -> pd.DataFrame:
    """Assign each session to the instructor or the TA."""
    assigned = []
    for course, group in sessions.groupby("course", sort=False):
        group = group.sort_values("date").reset_index(drop=True)
        expected = INSTRUCTOR_SESSIONS + ASSISTANT_SESSIONS
        if len(group) != expected:
            raise ValueError(f"{course}: expected {expected} sessions, found {len(group)}")

        owner = ["instructor"] * len(group)
        notes = [""] * len(group)

        # Rule 2: hand every clashing session to the TA
        for i, date in enumerate(group["date"]):
            if date in conflicts:
                owner[i] = "assistant"
                notes[i] = f"Cruce PUCP: {conflicts[date]}"

        # Rule 3: top up the TA's hours with the last sessions of the course
        for i in reversed(range(len(group))):
            if owner.count("assistant") >= ASSISTANT_SESSIONS:
                break
            if owner[i] == "instructor":
                owner[i] = "assistant"
                notes[i] = f"Balance de horas del JP ({ASSISTANT_SESSIONS * SESSION_HOURS} h)"

        group["owner"] = owner
        group["note"] = notes
        group["session_number"] = range(1, len(group) + 1)
        assigned.append(group)

    result = pd.concat(assigned, ignore_index=True)

    # Sanity checks: the split must hit 18/12 and leave no clash on the instructor
    for course, group in result.groupby("course"):
        hours = group.groupby("owner").size() * SESSION_HOURS
        assert hours["instructor"] == INSTRUCTOR_SESSIONS * SESSION_HOURS, course
        assert hours["assistant"] == ASSISTANT_SESSIONS * SESSION_HOURS, course
    clashing = result[(result["owner"] == "instructor") & (result["date"].isin(conflicts))]
    assert clashing.empty, f"instructor still double-booked on {list(clashing['date'])}"

    return result


def header(cell, text, background=BLUE, color="FFFFFF", size=11) -> None:
    cell.value = text
    cell.font = Font(bold=True, color=color, size=size)
    cell.fill = PatternFill("solid", fgColor=background)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BOX


def build_summary(sheet, assigned: pd.DataFrame) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["B2"] = "PEU-CD 2026 — ESCUELA NACIONAL DE ESTADÍSTICA E INFORMÁTICA (ENEI - INEI)"
    sheet["B2"].font = Font(bold=True, size=14, color=BLUE)
    sheet["B3"] = "Módulo II: Machine Learning — Distribución de horas Docente / Jefe de Prácticas"
    sheet["B3"].font = Font(bold=True, size=11, color="404040")

    columns = ["Curso", "Modalidad", "Responsable", "Rol", "Horas", "N° sesiones",
               "Fecha inicio", "Fecha término", "Frecuencia"]
    for offset, title in enumerate(columns):
        header(sheet.cell(row=6, column=2 + offset), title)

    rows = [
        ("II.1 Machine Learning I", "ML I", "instructor", "Docente", INSTRUCTOR),
        ("Prácticas de Machine Learning I", "ML I", "assistant", "Jefe de Prácticas", ASSISTANT),
        ("II.2 Machine Learning II", "ML II", "instructor", "Docente", INSTRUCTOR),
        ("Prácticas de Machine Learning II", "ML II", "assistant", "Jefe de Prácticas", ASSISTANT),
    ]
    row_index = 7
    for name, course, owner, role, person in rows:
        group = assigned[(assigned["course"] == course) & (assigned["owner"] == owner)]
        weekdays = sorted({d.weekday() for d in group["date"]})
        frequency = ", ".join(WEEKDAYS_ES[d] for d in weekdays) + " de 19:00 a 22:00 h"
        values = [name, "Virtual", person, role, len(group) * SESSION_HOURS, len(group),
                  min(group["date"]).strftime("%d/%m/%Y"),
                  max(group["date"]).strftime("%d/%m/%Y"), frequency]
        for offset, value in enumerate(values):
            cell = sheet.cell(row=row_index, column=2 + offset, value=value)
            cell.border = BOX
            cell.alignment = Alignment(
                horizontal="left" if offset in (0, 1, 2, 3, 8) else "center",
                vertical="center", wrap_text=offset in (0, 8))
            cell.fill = PatternFill("solid", fgColor=GREEN if owner == "instructor" else ORANGE)
        sheet.cell(row=row_index, column=2).font = Font(bold=True)
        row_index += 1

    for course in ("ML I", "ML II"):
        group = assigned[assigned["course"] == course]
        label = f"TOTAL Machine Learning {course.split()[1]}"
        sheet.cell(row=row_index, column=2, value=label).font = Font(bold=True)
        sheet.cell(row=row_index, column=6, value=len(group) * SESSION_HOURS).font = Font(bold=True)
        sheet.cell(row=row_index, column=7, value=len(group)).font = Font(bold=True)
        for column in range(2, 11):
            cell = sheet.cell(row=row_index, column=column)
            cell.border = BOX
            cell.fill = PatternFill("solid", fgColor=BLUE_LIGHT)
            if column in (6, 7):
                cell.alignment = Alignment(horizontal="center")
        row_index += 1

    notes = [
        "Notas",
        "1. Cada sesión dura 3 horas efectivas (19:00 a 22:00 h). La franja 18:00 a 19:00 h "
        "corresponde a charlas de las Direcciones Técnicas del INEI, lecturas y Proyecto Integrador.",
        "2. Cada curso tiene 30 horas: 18 horas a cargo del docente (6 sesiones) y 12 horas a cargo "
        "del jefe de prácticas (4 sesiones), igual que en la edición 2025.",
        "3. RESTRICCIÓN: los miércoles de 19:00 a 22:00 h el docente dicta Python / Python intermedio "
        "en la Diplomatura PUCP 2026 (calendario oficial, no modificable). Por eso TODAS las sesiones "
        "de miércoles del PEU-CD las cubre el jefe de prácticas.",
        "4. La sesión del lunes 26/10 se asigna al jefe de prácticas para completar sus 12 horas en "
        "ML II (los miércoles disponibles solo suman 9 h).",
        "5. Los sábados de 08:00 a 09:30 h el docente dicta en la PUCP; no hay cruce con el horario "
        "del PEU-CD (18:00 a 22:00 h).",
    ]
    row_index += 2
    for offset, note in enumerate(notes):
        cell = sheet.cell(row=row_index + offset, column=2, value=note)
        cell.font = Font(bold=offset == 0, size=10, color=BLUE if offset == 0 else "404040")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=row_index + offset, start_column=2,
                          end_row=row_index + offset, end_column=10)
        sheet.row_dimensions[row_index + offset].height = 18 if offset == 0 else 28

    for column, width in zip("BCDEFGHIJ", [38, 14, 20, 20, 9, 12, 14, 14, 34]):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[6].height = 30


def build_calendar(sheet, assigned: pd.DataFrame, enei: pd.DataFrame,
                   conflicts: dict[dt.date, str]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["B2"] = "HORARIO PEU-CD 2026 — MACHINE LEARNING I y II"
    sheet["B2"].font = Font(bold=True, size=14, color=BLUE)
    sheet["B3"] = f"Docente: {INSTRUCTOR}  ·  Jefe de Prácticas: {ASSISTANT}"
    sheet["B3"].font = Font(bold=True, size=11, color="404040")

    legend = [(f"Clase — {INSTRUCTOR}", GREEN, GREEN_TEXT),
              (f"Práctica — {ASSISTANT}", ORANGE, ORANGE_TEXT),
              ("Otro curso del programa", GREY, GREY_TEXT),
              ("Cruce con Diplomatura PUCP (no movible)", YELLOW, ORANGE_TEXT)]
    for index, (text, background, color) in enumerate(legend):
        cell = sheet.cell(row=5, column=2 + index * 2, value=text)
        cell.fill = PatternFill("solid", fgColor=background)
        cell.font = Font(bold=True, size=9, color=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
        sheet.merge_cells(start_row=5, start_column=2 + index * 2,
                          end_row=5, end_column=3 + index * 2)
    sheet.row_dimensions[5].height = 26

    header(sheet.cell(row=7, column=2), "MES")
    header(sheet.cell(row=7, column=3), "Horario")
    for offset, day in enumerate(WEEKDAYS_ES):
        header(sheet.cell(row=7, column=4 + offset), day)

    by_date = {row["date"]: row for _, row in assigned.iterrows()}
    other = {row["date"]: COURSE_LABELS.get(row["course"], row["course"])
             for _, row in enei.iterrows() if row["course"] not in COURSES}
    holidays = {d for d, label in other.items() if label.startswith("FERIADO")}

    first, last = assigned["date"].min(), assigned["date"].max()
    week = first - dt.timedelta(days=first.weekday())
    last_week = last - dt.timedelta(days=last.weekday())
    row_index, current_month = 8, None

    while week <= last_week:
        days = [week + dt.timedelta(days=k) for k in range(5)]

        for offset, day in enumerate(days):
            cell = sheet.cell(row=row_index, column=4 + offset,
                              value=f"{day.day:02d} de {MONTHS_ES[day.month]}")
            cell.font = Font(bold=True, size=10)
            cell.border = BOX
            cell.fill = PatternFill("solid", fgColor=BLUE_LIGHT)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row=row_index, column=2).border = BOX
        sheet.cell(row=row_index, column=3).border = BOX
        if days[0].month != current_month:
            current_month = days[0].month
            cell = sheet.cell(row=row_index, column=2, value=MONTHS_ES[current_month].upper())
            cell.font = Font(bold=True, size=11, color=BLUE)
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center", vertical="center")
        row_index += 1

        cell = sheet.cell(row=row_index, column=3, value="18:00 a 19:00")
        cell.font = Font(size=9)
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row=row_index, column=2).border = BOX
        for offset, day in enumerate(days):
            cell = sheet.cell(row=row_index, column=4 + offset,
                              value="" if day in holidays else
                              "Charlas DDTT del INEI / lecturas / Proyecto Integrador")
            cell.font = Font(size=8, color=GREY_TEXT, italic=True)
            cell.border = BOX
            cell.fill = PatternFill("solid", fgColor=GREY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[row_index].height = 26
        row_index += 1

        cell = sheet.cell(row=row_index, column=3, value="19:00 a 22:00")
        cell.font = Font(size=9, bold=True)
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row=row_index, column=2).border = BOX
        for offset, day in enumerate(days):
            cell = sheet.cell(row=row_index, column=4 + offset)
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            session = by_date.get(day)
            if session is not None:
                number = session["session_number"]
                if session["owner"] == "instructor":
                    cell.value = f"{session['course']} · Sesión {number}\\nCLASE — {INSTRUCTOR}"
                    cell.fill = PatternFill("solid", fgColor=GREEN)
                    cell.font = Font(size=9, bold=True, color=GREEN_TEXT)
                else:
                    extra = "\\n(PUCP: Python)" if day in conflicts else ""
                    cell.value = f"{session['course']} · Sesión {number}\\nPRÁCTICA — {ASSISTANT}{extra}"
                    cell.fill = PatternFill("solid", fgColor=ORANGE)
                    cell.font = Font(size=9, bold=True, color=ORANGE_TEXT)
            else:
                cell.value = other.get(day, "")
                cell.font = Font(size=8, color=GREY_TEXT)
                cell.fill = PatternFill("solid", fgColor=GREY)
        sheet.row_dimensions[row_index].height = 40
        row_index += 2
        week += dt.timedelta(days=7)

    sheet.column_dimensions["B"].width = 13
    sheet.column_dimensions["C"].width = 14
    for column in "DEFGH":
        sheet.column_dimensions[column].width = 27
    sheet.row_dimensions[7].height = 22


def build_detail(sheet, assigned: pd.DataFrame) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["B2"] = "DETALLE DE SESIONES — ML I y ML II (PEU-CD 2026)"
    sheet["B2"].font = Font(bold=True, size=14, color=BLUE)

    columns = ["N°", "Curso", "Fecha", "Día", "Horario", "Horas",
               "Tipo de sesión", "Responsable", "Observación"]
    for offset, title in enumerate(columns):
        header(sheet.cell(row=4, column=2 + offset), title)

    row_index = 5
    for _, session in assigned.iterrows():
        is_instructor = session["owner"] == "instructor"
        values = [session["session_number"],
                  f"Machine Learning {session['course'].split()[1]}",
                  session["date"].strftime("%d/%m/%Y"),
                  WEEKDAYS_ES[session["date"].weekday()],
                  "19:00 - 22:00", SESSION_HOURS,
                  "Clase" if is_instructor else "Práctica",
                  INSTRUCTOR if is_instructor else ASSISTANT,
                  session["note"]]
        for offset, value in enumerate(values):
            cell = sheet.cell(row=row_index, column=2 + offset, value=value)
            cell.border = BOX
            cell.alignment = Alignment(
                horizontal="center" if offset in (0, 2, 3, 4, 5, 6) else "left",
                vertical="center", wrap_text=offset == 8)
            cell.font = Font(size=10)
            cell.fill = PatternFill("solid", fgColor=GREEN if is_instructor else ORANGE)
            if offset == 8 and str(value).startswith("Cruce"):
                cell.fill = PatternFill("solid", fgColor=YELLOW)
                cell.font = Font(size=9, bold=True, color=ORANGE_TEXT)
        sheet.row_dimensions[row_index].height = 26
        row_index += 1

    row_index += 1
    header(sheet.cell(row=row_index, column=2), "Verificación de horas")
    sheet.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=5)
    header(sheet.cell(row=row_index, column=6), "Sesiones")
    header(sheet.cell(row=row_index, column=7), "Horas")
    row_index += 1

    for course in ("ML I", "ML II"):
        for owner, person, role in (("instructor", INSTRUCTOR, "docente"),
                                    ("assistant", ASSISTANT, "jefe de prácticas")):
            group = assigned[(assigned["course"] == course) & (assigned["owner"] == owner)]
            label = f"Machine Learning {course.split()[1]} — {person} ({role})"
            cell = sheet.cell(row=row_index, column=2, value=label)
            cell.font = Font(size=10, bold=True)
            sheet.merge_cells(start_row=row_index, start_column=2,
                              end_row=row_index, end_column=5)
            sheet.cell(row=row_index, column=6, value=len(group))
            sheet.cell(row=row_index, column=7, value=len(group) * SESSION_HOURS)
            for column in range(2, 8):
                cell = sheet.cell(row=row_index, column=column)
                cell.border = BOX
                cell.fill = PatternFill("solid",
                                        fgColor=GREEN if owner == "instructor" else ORANGE)
                if column in (6, 7):
                    cell.alignment = Alignment(horizontal="center")
                    cell.font = Font(size=10)
            row_index += 1

    for column, width in zip("BCDEFGHIJ", [6, 22, 12, 12, 14, 8, 15, 20, 38]):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[4].height = 30
    sheet.freeze_panes = "B5"


def main() -> None:
    sessions, conflicts = load_sessions()
    assigned = assign(sessions, conflicts)
    enei = pd.read_csv(PROCESSED / "enei_sessions_2026.csv", parse_dates=["date"])
    enei["date"] = enei["date"].dt.date

    workbook = openpyxl.Workbook()
    build_summary(workbook.active, assigned)
    workbook.active.title = "RESUMEN"
    build_calendar(workbook.create_sheet("HORARIO DETALLE ML-I ML-II"), assigned, enei, conflicts)
    build_detail(workbook.create_sheet("DETALLE SESIONES"), assigned)

    OUTPUT.mkdir(parents=True, exist_ok=True)
    destination = OUTPUT / "horario_peu_cd_2026_ml1_ml2.xlsx"
    workbook.save(destination)
    assigned.to_csv(PROCESSED / "ml_teaching_split_2026.csv", index=False)

    print(f"Saved {destination.relative_to(ROOT.parent)}\n")
    for course, group in assigned.groupby("course"):
        instructor_hours = (group["owner"] == "instructor").sum() * SESSION_HOURS
        assistant_hours = (group["owner"] == "assistant").sum() * SESSION_HOURS
        print(f"{course:6s} {len(group)} sessions = {len(group) * SESSION_HOURS} h | "
              f"{INSTRUCTOR}: {instructor_hours} h | {ASSISTANT}: {assistant_hours} h")
    print(f"\nWednesday clashes covered by the TA: "
          f"{sum(d in conflicts for d in assigned['date'])}")


if __name__ == "__main__":
    main()
