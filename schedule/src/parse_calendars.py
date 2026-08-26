"""Parse the two source calendars and write tidy CSV files.

Inputs (schedule/data/raw/):
    enei_peu_cd_2026_program_calendar.xlsx  ENEI-INEI program calendar for PEU-CD 2026
    pucp_diploma_2026_official_calendar.pdf PUCP Diploma 2026 official calendar

Outputs (schedule/data/processed/):
    enei_sessions_2026.csv    every ENEI evening session, one row per date
    pucp_conflicts_2026.csv   evening classes taught by Alexander Quispe at PUCP
"""

import datetime as dt
import re
from pathlib import Path

import pandas as pd
import openpyxl
import pdfplumber

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
PROCESSED = ROOT / "data" / "processed"

INSTRUCTOR = "Alexander Quispe"

# Column x-offsets of the weekday headers in the PUCP PDF calendar
PDF_DAY_COLUMNS = [25.7, 129.2, 232.3, 335.4, 438.5, 541.6, 644.7]
WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def parse_enei_calendar() -> pd.DataFrame:
    """Read the ENEI program calendar (one merged cell per day) into tidy rows."""
    workbook = openpyxl.load_workbook(RAW / "enei_peu_cd_2026_program_calendar.xlsx", data_only=True)
    sheet = workbook["HORARIO DETALLADO"]

    rows = []
    for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row, min_col=2, max_col=6):
        for cell in row:
            if not cell.value:
                continue
            text = str(cell.value).strip()
            match = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})\s*(.*)", text)
            if not match:
                continue

            day, month, year, course = (
                int(match.group(1)),
                int(match.group(2)),
                int(match.group(3)),
                match.group(4).strip(),
            )
            weekday_index = cell.column - 2  # column B is Monday

            # The whole calendar is 2026; the file has a couple of year typos.
            year = 2026
            try:
                date = dt.date(year, month, day)
            except ValueError:
                continue

            # The source file has a few typos in the month/year. The grid position
            # is authoritative, so shift the month until the weekday matches.
            if date.weekday() != weekday_index:
                for candidate_month in (month - 1, month + 1):
                    try:
                        candidate = dt.date(2026, candidate_month, day)
                    except ValueError:
                        continue
                    if candidate.weekday() == weekday_index:
                        date = candidate
                        break

            notes = " | ".join(line.strip() for line in text.split("\n")[1:] if line.strip())
            rows.append({"date": date, "weekday": WEEKDAYS[date.weekday()],
                         "course": course, "notes": notes})

    return pd.DataFrame(rows).sort_values("date").reset_index(drop=True)


def _pdf_column(x0: float) -> int:
    """Map a word's x-offset to the weekday column it belongs to."""
    column = 0
    for index, boundary in enumerate(PDF_DAY_COLUMNS):
        if x0 >= boundary - 8:
            column = index
    return column


def parse_pucp_calendar() -> pd.DataFrame:
    """Extract the PUCP evening classes taught by the instructor.

    The PDF is a month grid. Words are bucketed into (row band, weekday column);
    the bands where most columns hold a bare number are the day-number rows, and
    everything between two such rows belongs to that week's cells. A week can
    spill over a page break, so the last day-number row carries over.
    """
    cells = []  # (day, weekday_index, text)
    pending: dict[int, int] | None = None  # weekday column -> day number

    with pdfplumber.open(RAW / "pucp_diploma_2026_official_calendar.pdf") as pdf:
        for page in pdf.pages:
            bands: dict[int, dict[int, list]] = {}
            for word in page.extract_words():
                band = round(word["top"] / 6)
                bands.setdefault(band, {}).setdefault(_pdf_column(word["x0"]), []).append(word)

            buffer: dict[int, list[str]] = {}

            def flush() -> None:
                if pending is None:
                    return
                for column, chunks in buffer.items():
                    if column in pending:
                        cells.append((pending[column], column, " ".join(chunks)))

            for band in sorted(bands):
                columns = bands[band]
                texts = {
                    column: " ".join(w["text"] for w in sorted(words, key=lambda w: w["x0"]))
                    for column, words in columns.items()
                }
                is_day_row = (
                    len(texts) >= 5
                    and all(re.fullmatch(r"\d{1,2}", t.strip()) for t in texts.values())
                )
                if is_day_row:
                    flush()
                    buffer = {}
                    pending = {column: int(t) for column, t in texts.items()}
                    continue
                for column, text in texts.items():
                    buffer.setdefault(column, []).append(text)
            flush()

    rows = [
        {"weekday": WEEKDAYS[column], "day": day, "raw_text": " ".join(text.split())}
        for day, column, text in cells
        if "Quispe" in text
    ]
    conflicts = pd.DataFrame(rows).drop_duplicates()

    # Keep the evening classes only: those are the ones that clash with the ENEI
    # schedule (18:00-22:00). Saturday and Sunday classes start at 08:00.
    conflicts = conflicts[conflicts["weekday"] == "Wednesday"].copy()
    conflicts["date"] = conflicts["day"].map(_resolve_wednesday)
    conflicts = conflicts.dropna(subset=["date"]).sort_values("date")
    conflicts["label"] = conflicts["raw_text"].map(_clean_label)

    return conflicts[["date", "weekday", "label", "raw_text"]].reset_index(drop=True)


def _clean_label(text: str) -> str:
    """Rebuild a readable class label from the scrambled PDF text.

    Words from adjacent lines interleave when the PDF is flattened, so the
    class number, the course name and the time range are pulled out separately.
    """
    number = re.search(r"\b(\d{1,2})\s*-", text)
    course = "Python intermedio" if "intermedio" in text.lower() else "Python"
    label = f"Clase {number.group(1)} - {course}" if number else course

    hours = re.search(r"(\d{1,2})(?::\d{2})?\s*a\s*(\d{1,2})(?::(\d{2}))?\s*[AP]M", text)
    if hours:
        start, end, minutes = int(hours.group(1)), int(hours.group(2)), hours.group(3) or "00"
        # The calendar writes evening hours in 12-hour format (and one "AM" typo)
        label += f" ({start + 12}:00-{end + 12}:{minutes})"
    return label


def _resolve_wednesday(day: int):
    """Map a day-of-month to a full date within the teaching term.

    Between September and November 2026 every day-of-month that falls on a
    Wednesday is unique, so the mapping is unambiguous.
    """
    matches = []
    date = dt.date(2026, 9, 1)
    while date <= dt.date(2026, 11, 30):
        if date.weekday() == 2 and date.day == day:
            matches.append(date)
        date += dt.timedelta(days=1)
    return matches[0] if len(matches) == 1 else None


def main() -> None:
    PROCESSED.mkdir(parents=True, exist_ok=True)

    enei = parse_enei_calendar()
    enei.to_csv(PROCESSED / "enei_sessions_2026.csv", index=False)
    print(f"enei_sessions_2026.csv        {len(enei)} rows "
          f"({enei['date'].min()} to {enei['date'].max()})")

    ml_sessions = enei[enei["course"].str.contains("LEARNING", na=False)]
    print("\nMachine Learning sessions in the ENEI calendar:")
    for course, group in ml_sessions.groupby("course"):
        print(f"  {course:22s} {len(group):2d} sessions  "
              f"{group['date'].min()} to {group['date'].max()}")

    pucp = parse_pucp_calendar()
    pucp.to_csv(PROCESSED / "pucp_conflicts_2026.csv", index=False)
    print(f"\npucp_conflicts_2026.csv       {len(pucp)} rows")
    print(f"\n{INSTRUCTOR} teaches Wednesday evenings at PUCP (cannot be moved):")
    for _, row in pucp.iterrows():
        print(f"  {row['date']}  {row['label']}")


if __name__ == "__main__":
    main()
